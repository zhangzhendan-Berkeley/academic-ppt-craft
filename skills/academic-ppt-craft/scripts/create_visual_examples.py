"""Create a compact editable style exercise, without any research or confidential assets."""
from pathlib import Path
import json, hashlib
from pptx import Presentation
from pptx.util import Inches,Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN

root=Path(__file__).resolve().parents[1]
p=Presentation();p.slide_width=Inches(13.333);p.slide_height=Inches(7.5)
B='0B4BA3';R='C00000';K='171717'
def box(sl,x,y,w,h,fill=None,stroke=None,width=2):
 s=sl.shapes.add_shape(MSO_SHAPE.RECTANGLE,Inches(x),Inches(y),Inches(w),Inches(h))
 if fill:s.fill.solid();s.fill.fore_color.rgb=RGBColor.from_string(fill)
 else:s.fill.background()
 if stroke:s.line.color.rgb=RGBColor.from_string(stroke);s.line.width=Pt(width)
 else:s.line.fill.background()
 return s
def text(sl,t,x,y,w,h,color=K,size=24,fill=None):
 s=sl.shapes.add_textbox(Inches(x),Inches(y),Inches(w),Inches(h));s.text=t
 if fill:s.fill.solid();s.fill.fore_color.rgb=RGBColor.from_string(fill)
 for para in s.text_frame.paragraphs:
  para.font.name='Microsoft YaHei';para.font.size=Pt(size);para.font.color.rgb=RGBColor.from_string(color);para.font.bold=True
 return s
def placeholder(sl,x,y,w,h,t='实验画面'):
 box(sl,x,y,w,h,'EDF1F5');text(sl,t,x+.2,y+h/2-.3,w-.4,.6,'6B7789')

titles=['白底图名压住粗边框','并列条件与共同外框','角标与局部说明','侧边色线联系解释','视频画面与色签','部件、过程与系统']
descs=['白底只遮住图名背后的边线；图、框、文字均可编辑。','同一任务的两个条件，保留相同尺度与对齐关系。','多个同类小图可以少用边框，靠角标和短标签组织。','图侧说明对准图内对象，边线不绕整页一圈。','色签在画面外侧，避免遮住工件、手指和实验数据。','箭头只表达真实先后或连接；并列内容不强画流程。']
for i,(title,desc) in enumerate(zip(titles,descs)):
 sl=p.slides.add_slide(p.slide_layouts[6]);text(sl,title,.5,.3,12.2,.7,B,30);text(sl,desc,.5,1.1,12.2,.7,size=21)
 if i==0:
  placeholder(sl,2.5,2.15,8.3,3.8);box(sl,2.4,2.05,8.5,4.3,stroke=B,width=2.5)
  text(sl,'关键实验',5.15,6.02,3.0,.62,B,26,'FFFFFF')
 elif i==1:
  box(sl,.8,2.2,11.7,4.15,stroke=B,width=2)
  for j,(name,col) in enumerate([('条件 A',B),('条件 B',R)]):
   placeholder(sl,1.1+j*5.8,2.6,5.15,2.95);text(sl,name,2.55+j*5.8,5.95,2.25,.62,col,25,'FFFFFF')
 elif i==2:
  for j in range(3):
   x=.8+j*4.2;placeholder(sl,x,2.65,3.45,2.55)
   box(sl,x-.07,2.55,.06,.52,B);box(sl,x-.07,2.55,.52,.06,B)
   text(sl,['接触前','接触建立','接触变化'][j],x+.4,5.4,3,.65,B)
 elif i==3:
  placeholder(sl,.9,2.4,7.0,3.8);box(sl,.75,2.4,.06,3.8,R)
  text(sl,'部件与观测',8.5,2.7,4.0,.6,B);text(sl,'说明对应图中的具体对象。',8.5,3.45,4.0,1.1,size=22)
  text(sl,'实验结果',8.5,4.8,4,.6,R);text(sl,'短句说明条件与结论。',8.5,5.5,4.0,.75,size=22)
 elif i==4:
  placeholder(sl,2.4,2.15,8.5,3.9,'视频区域（此附件不含媒体）');box(sl,2.36,2.11,8.58,3.98,stroke=B,width=1.5)
  text(sl,'任务与接触过程',4.8,6.18,3.9,.65,'FFFFFF',24,B)
 else:
  for j,t in enumerate(['器件','接触过程','系统观测']):
   x=.8+j*4.25;placeholder(sl,x,2.5,3.2,3.0,t);text(sl,t,x+.35,5.7,2.8,.7,B)
   if j<2:
    arrow=sl.shapes.add_shape(MSO_SHAPE.RIGHT_ARROW,Inches(x+3.4),Inches(3.75),Inches(.58),Inches(.5));arrow.fill.solid();arrow.fill.fore_color.rgb=RGBColor.from_string(R);arrow.line.fill.background()
 sl.notes_slide.notes_text_frame.text='这是图文组合练习文件；灰色区域为中性占位，不是课程实验或研究结果。原生矩形、文本和箭头可编辑。'
out=root/'assets/visual-combinations.pptx';p.save(out)
f=root/'assets/attachment-manifest.json';m=json.loads(f.read_text('utf-8'))
m['attachments']=[x for x in m['attachments'] if x['path']!='assets/visual-combinations.pptx']
m['attachments'].append(dict(path='assets/visual-combinations.pptx',source_name='Agent-created native decoration exercises',kind='editable_style_examples',original_pages=list(range(1,7)),sha256=hashlib.sha256(out.read_bytes()).hexdigest(),bytes=out.stat().st_size,scope='Neutral diagram placeholders only; no confidential research content.',media=[]))
f.write_text(json.dumps(m,ensure_ascii=False,indent=2),'utf-8')
print(out)
